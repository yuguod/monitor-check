# -*- coding: utf-8 -*-
"""端到端验证: 连跑两次 exe, 检查 Excel 是否新建 + 追加(两行), 并截屏看提示"""
import os
import subprocess
import time

from openpyxl import load_workbook
from PIL import ImageGrab

BASE = r"C:\Users\yuguo\WorkBuddy\2026-09-03-20-07-00\monitor-check"
EXE = BASE + r"\MonitorCheck.exe"
XLSX = BASE + r"\显示器检测记录.xlsx"


def show(tag):
    if not os.path.exists(XLSX):
        print(tag, "-> xlsx 不存在")
        return
    ws = load_workbook(XLSX).active
    print(tag, "->", XLSX, "行数(含表头):", ws.max_row)
    for r in ws.iter_rows(values_only=True):
        print("   ", r)


if os.path.exists(XLSX):
    os.remove(XLSX)

# 第一次运行
p = subprocess.Popen([EXE])
time.sleep(9.0)
ImageGrab.grab().resize((960, 600)).save(BASE + r"\python\shot_x1.png")
p.kill()
time.sleep(1.0)
show("run1")

# 第二次运行 → 应追加
time.sleep(2.0)
p = subprocess.Popen([EXE])
time.sleep(9.0)
ImageGrab.grab().resize((960, 600)).save(BASE + r"\python\shot_x2.png")
p.kill()
time.sleep(1.0)
show("run2")
