# -*- coding: utf-8 -*-
"""
MonitorCheck v3.1 - 显示器检测工具 (原生 Windows 版)
- 显示显示器型号信息 (WMI 主通道 + 注册表 EDID 回退)
- 19 种测试画面循环: 纯色 / 灰阶 / 彩条 / SMPTE / 渐变 / 棋盘 / 网格
- 空格 / 左键: 下一画面    左方向键: 上一画面
- 右键双击 / ESC: 退出     P: 暂停/恢复自动循环
- R: 开启/关闭检测记录(可选, 默认关闭, 选择会记住)
- 运行期间阻止系统休眠与熄屏; 测试画面中上部常驻信息栏, 每 10 秒刷新
纯 tkinter 实现, 不依赖 IE/mshta, 无任何安全提示。
"""
import ctypes
import os
import subprocess
import sys
import threading
import time
import tkinter as tk
import tkinter.font as tkfont

APP_VERSION = "3.1"
APP_NAME = "MonitorCheck 显示器检测"
APP_TITLE = APP_NAME + " v" + APP_VERSION
AUTO_MS = 2000            # 自动循环间隔
REFRESH_MS = 10000        # 信息栏刷新间隔


# ---------------------------------------------------------------- 错误日志
def log_error(context, exc):
    """把异常写入 %TEMP%\\MonitorCheck_err.log (静默失败, 不影响主流程)"""
    try:
        import os
        import traceback
        path = os.path.join(os.environ.get("TEMP", "."), "MonitorCheck_err.log")
        with open(path, "a", encoding="utf-8") as f:
            f.write("\n[%s] %s: %s\n" % (time.strftime("%Y-%m-%d %H:%M:%S"),
                                        context, exc))
            f.write(traceback.format_exc())
    except Exception:
        pass


# ---------------------------------------------------------------- 显示器信息
def _decode_wmi_bytes(raw):
    """WmiMonitorID 的字节数组(每字节后跟0)转字符串"""
    try:
        b = bytes(int(x) for x in raw if int(x) != 0)
        return b.decode("utf-8", "ignore").strip()
    except Exception:
        return ""


def get_monitors_wmi():
    """主通道: PowerShell 查询 WmiMonitorID / WmiMonitorBasicDisplayParams"""
    monitors = []
    try:
        ps = (
            "Get-CimInstance -Namespace root\\wmi -ClassName WmiMonitorID | "
            "ForEach-Object { "
            "$n=($_.UserFriendlyName | Where-Object {$_ -ne 0}); "
            "$m=($_.ManufacturerName | Where-Object {$_ -ne 0}); "
            "$s=($_.SerialNumberID | Where-Object {$_ -ne 0}); "
            "[pscustomobject]@{M=(-join [char[]]$m); N=(-join [char[]]$n); "
            "S=(-join [char[]]$s); Year=$_.YearOfManufacture; "
            "Active=$_.Active} } | ConvertTo-Json -Compress"
        )
        out = subprocess.run(
            ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", ps],
            capture_output=True, timeout=10,
            creationflags=subprocess.CREATE_NO_WINDOW,
        )
        text = out.stdout.decode("utf-8", "ignore").strip()
        import json
        if text:
            data = json.loads(text)
            if isinstance(data, dict):
                data = [data]
            for d in data:
                if not d.get("Active"):
                    continue
                mfr = (d.get("M") or "").strip()
                name = (d.get("N") or "").strip()
                ser = (d.get("S") or "").strip()
                year = d.get("Year") or ""
                monitors.append({"mfr": mfr, "name": name,
                                 "serial": ser, "year": str(year)})
    except Exception:
        pass
    return monitors


MFR_MAP = {
    "LEN": "联想", "DEL": "戴尔", "HPH": "惠普", "AOC": "AOC", "ACI": "华硕",
    "GSM": "LG", "SAM": "三星", "BNQ": "明基", "PHL": "飞利浦", "MSI": "微星",
    "APP": "苹果", "NEC": "NEC", "HWP": "惠普", "YMH": "雅马哈",
}


def _parse_edid(edid):
    """解析 128 字节 EDID: 型号(0xFE/0xFC) / 序列号(0xFF / 12-15) / 出厂年月(16,17)"""
    info = {"name": "", "name_fc": "", "serial_str": "", "serial_num": 0,
            "year": 0, "week": 0, "ym": ""}
    if len(edid) < 128:
        return info
    info["week"] = edid[16]
    info["year"] = edid[17] + 1990
    if info["year"] >= 1990:
        if 1 <= info["week"] <= 53:
            month = min(12, max(1, round(info["week"] / 4.33)))
            info["ym"] = "%d年%02d月" % (info["year"], month)
        else:
            info["ym"] = "%d年" % info["year"]
    info["serial_num"] = int.from_bytes(edid[12:16], "little")
    for k in range(4):
        off = 54 + k * 18
        if off + 18 > len(edid):
            break
        if edid[off] == 0 and edid[off + 1] == 0:      # 描述符
            tag = edid[off + 3]
            val = edid[off + 5:off + 18].decode("ascii", "ignore")
            val = val.replace("\n", " ").replace("\r", " ").strip()
            if not val:
                continue
            if tag == 0xFE and not info["name"]:
                info["name"] = val                    # 面板名(笔记本常用)
            elif tag == 0xFC and not info["name_fc"]:
                info["name_fc"] = val                 # 显示器型号
            elif tag == 0xFF and not info["serial_str"]:
                info["serial_str"] = val              # 序列号字符串
        # else: 时序描述符, 跳过
    return info


def get_monitors_edid():
    """回退通道: 注册表 EDID 读取 (型号 / 序列号 / 出厂年月)"""
    results = []
    try:
        import winreg
        base = r"SYSTEM\CurrentControlSet\Enum\DISPLAY"
        key = winreg.OpenKey(winreg.HKEY_LOCAL_MACHINE, base)
        i = 0
        while True:
            try:
                mfr = winreg.EnumKey(key, i); i += 1
            except OSError:
                break
            try:
                mkey = winreg.OpenKey(key, mfr)
            except OSError:
                continue
            j = 0
            while True:
                try:
                    dev = winreg.EnumKey(mkey, j); j += 1
                except OSError:
                    break
                try:
                    dkey = winreg.OpenKey(mkey, dev + r"\Device Parameters")
                    edid, _ = winreg.QueryValueEx(dkey, "EDID")
                    dkey.Close()
                except OSError:
                    continue
                edid = bytes(edid)
                if len(edid) < 128:
                    continue
                p = _parse_edid(edid)
                name = p["name"] or p["name_fc"]
                serial = p["serial_str"] or (
                    str(p["serial_num"]) if p["serial_num"] else "")
                results.append({"mfr": mfr, "name": name, "serial": serial,
                                "year": str(p["year"]) if p["year"] else "",
                                "ym": p["ym"]})
                mkey.Close()
        key.Close()
    except Exception:
        pass
    return results


def merge_monitors():
    wmi = get_monitors_wmi()
    edid = get_monitors_edid()
    merged = []
    for w in wmi:
        name = w["name"]
        if not name:
            # WMI 没给名字时, 用 EDID 的 0xFE 面板型号补
            # (注册表键名形如 LEN9123, WMI 厂商码是 LEN, 用前缀匹配)
            for e in edid:
                if e["name"] and e["mfr"].upper().startswith(w["mfr"].upper()):
                    name = e["name"]
                    break
        merged.append({**w, "name": name})
    if not merged:
        merged = edid
    return merged


def pretty_line(m):
    mfr = m.get("mfr", "")
    cn = MFR_MAP.get(mfr.upper(), mfr)
    name = m.get("name", "") or "(未知型号)"
    year = m.get("year", "")
    ser = m.get("serial", "")
    label = f"{cn} {name}" if cn and cn != mfr else f"{mfr} {name}"
    extra = []
    if year and year != "0":
        extra.append(f"{year} 年产")
    if ser and ser.strip("0 "):
        extra.append(f"SN: {ser}")
    return label + ("   " + " · ".join(extra) if extra else "")


def get_gpu_names():
    """读取显卡名称 (PowerShell, 只在启动/拓扑变化时调用一次)"""
    try:
        out = subprocess.run(
            ["powershell", "-NoProfile", "-Command",
             "Get-CimInstance Win32_VideoController | ForEach-Object { $_.Name }"],
            capture_output=True, timeout=10,
            creationflags=subprocess.CREATE_NO_WINDOW,
        )
        text = out.stdout.decode("utf-8", "ignore").strip()
        return [l.strip() for l in text.splitlines() if l.strip()]
    except Exception:
        return []


class _DISPLAY_DEVICEW(ctypes.Structure):
    _fields_ = [("cb", ctypes.c_ulong),
                ("DeviceName", ctypes.c_wchar * 32),
                ("DeviceString", ctypes.c_wchar * 128),
                ("StateFlags", ctypes.c_ulong),
                ("DeviceID", ctypes.c_wchar * 128),
                ("DeviceKey", ctypes.c_wchar * 128)]


class _DEVMODEW(ctypes.Structure):
    _fields_ = [
        ("dmDeviceName", ctypes.c_wchar * 32),
        ("dmSpecVersion", ctypes.c_ushort), ("dmDriverVersion", ctypes.c_ushort),
        ("dmSize", ctypes.c_ushort), ("dmDriverExtra", ctypes.c_ushort),
        ("dmFields", ctypes.c_ulong),
        ("dmOrientation", ctypes.c_short), ("dmPaperSize", ctypes.c_short),
        ("dmPaperLength", ctypes.c_short), ("dmPaperWidth", ctypes.c_short),
        ("dmScale", ctypes.c_short), ("dmCopies", ctypes.c_short),
        ("dmDefaultSource", ctypes.c_short), ("dmPrintQuality", ctypes.c_short),
        ("dmColor", ctypes.c_short), ("dmDuplex", ctypes.c_short),
        ("dmYResolution", ctypes.c_short), ("dmTTOption", ctypes.c_short),
        ("dmCollate", ctypes.c_short), ("dmFormName", ctypes.c_wchar * 32),
        ("dmLogPixels", ctypes.c_ushort),
        ("dmBitsPerPel", ctypes.c_ulong), ("dmPelsWidth", ctypes.c_ulong),
        ("dmPelsHeight", ctypes.c_ulong), ("dmDisplayFlags", ctypes.c_ulong),
        ("dmDisplayFrequency", ctypes.c_ulong),
        ("dmICMMethod", ctypes.c_ulong), ("dmICMIntent", ctypes.c_ulong),
        ("dmMediaType", ctypes.c_ulong), ("dmDitherType", ctypes.c_ulong),
        ("dmReserved1", ctypes.c_ulong), ("dmReserved2", ctypes.c_ulong),
        ("dmPanningWidth", ctypes.c_ulong), ("dmPanningHeight", ctypes.c_ulong),
    ]


def enum_active_displays():
    """轻量枚举活动显示器(纯 ctypes, 无子进程, 毫秒级)。
    返回 [{'pnp': 'LEN9123', 'w':2880, 'h':1800, 'hz':60, 'primary':True}]"""
    user32 = ctypes.windll.user32
    out = []
    try:
        user32.SetProcessDPIAware()
        pmw, pmh = user32.GetSystemMetrics(0), user32.GetSystemMetrics(1)
    except Exception:
        pmw = pmh = 0
    i = 0
    while True:
        dd = _DISPLAY_DEVICEW()
        dd.cb = ctypes.sizeof(dd)
        if not user32.EnumDisplayDevicesW(None, i, ctypes.byref(dd), 0):
            break
        i += 1
        if not (dd.StateFlags & 0x1):      # DISPLAY_DEVICE_ATTACHED_TO_DESKTOP
            continue
        mon = _DISPLAY_DEVICEW()
        mon.cb = ctypes.sizeof(mon)
        pnp = ""
        if user32.EnumDisplayDevicesW(dd.DeviceName, 0, ctypes.byref(mon), 0):
            parts = mon.DeviceID.split("\\")
            if len(parts) > 1:
                pnp = parts[1].upper()
        dm = _DEVMODEW()
        dm.dmSize = ctypes.sizeof(dm)
        w = h = hz = 0
        if user32.EnumDisplaySettingsW(dd.DeviceName, -1, ctypes.byref(dm)):
            w, h, hz = dm.dmPelsWidth, dm.dmPelsHeight, dm.dmDisplayFrequency
        if not pnp and not w:
            continue
        out.append({"pnp": pnp, "w": w, "h": h, "hz": hz,
                    "primary": (w == pmw and h == pmh)})
    return out


def edid_panel_map():
    """注册表 EDID 面板名映射 {PnP_ID: 面板型号}"""
    m = {}
    for e in get_monitors_edid():
        if e["name"]:
            m[e["mfr"].upper()] = e["name"]
    return m


def quick_monitor_lines(wmi_cache, edid_map):
    """快速组装显示信息行(每 10 秒刷新用, 不启动任何子进程)"""
    displays = enum_active_displays()
    lines = []
    for d in displays:
        pnp = d["pnp"]
        mfr = pnp[:3] if pnp else ""
        wmi = next((x for x in wmi_cache
                    if mfr and x["mfr"].upper().startswith(mfr)), None)
        name = edid_map.get(pnp) or (wmi["name"] if wmi else "") or "(未知型号)"
        cn = MFR_MAP.get(mfr, mfr)
        label = f"{cn} {name}" if cn and cn != mfr else f"{pnp} {name}"
        res = f"{d['w']}x{d['h']}" if d["w"] else ""
        if res and d["hz"]:
            res += f" @{d['hz']}Hz"
        seg = [x for x in ((res), (f"{wmi['year']}年" if wmi and wmi.get("year")
                                    and wmi["year"] != "0" else "")) if x]
        lines.append("● " + label + ("   " + " · ".join(seg) if seg else ""))
    if not lines:
        lines = ["● (未检测到活动显示器)"]
    ids = tuple(sorted(d["pnp"] or f"{d['w']}x{d['h']}" for d in displays))
    return lines, ids


# ---------------------------------------------------------------- 检测记录(xlsx)
try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    _HAS_OPENPYXL = True
except Exception:
    _HAS_OPENPYXL = False

REPORT_NAME = "显示器检测记录.xlsx"
HEADERS = ["检测时间", "品牌", "型号", "序列号", "出厂年月",
           "分辨率", "刷新率"]
COL_WIDTHS = [20, 12, 28, 24, 14, 14, 10]


def report_path():
    """记录文件放在 exe 同目录(打包后取 exe 所在目录, 而非临时解包目录)"""
    if getattr(sys, "frozen", False):
        base = os.path.dirname(sys.executable)
    else:
        base = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base, REPORT_NAME)


def fallback_report_path():
    """exe 目录不可写时(如放在 Program Files) → 依次退到桌面/用户目录/临时目录"""
    cands = [os.path.join(os.path.expanduser("~"), "Desktop"),
             os.path.expanduser("~"),
             os.environ.get("TEMP", "")]
    for c in cands:
        try:
            if c and os.path.isdir(c) and os.access(c, os.W_OK):
                return os.path.join(c, REPORT_NAME)
        except Exception:
            continue
    return os.path.join(os.environ.get("TEMP", "."), REPORT_NAME)


# ---------------------------------------------------------------- 配置(记录开关)
# 记录是可选项, 默认关闭; 用户的选择持久化在 %APPDATA%\MonitorCheck\config.ini
CONFIG_DIR = os.path.join(
    os.environ.get("APPDATA") or os.path.expanduser("~"), "MonitorCheck")
CONFIG_PATH = os.path.join(CONFIG_DIR, "config.ini")


def load_record_enabled():
    """读取上次选择; 文件不存在/损坏一律视为关闭"""
    try:
        import configparser
        cp = configparser.ConfigParser()
        if cp.read(CONFIG_PATH, encoding="utf-8") and cp.has_option(
                "record", "enabled"):
            return cp.getboolean("record", "enabled")
    except Exception as e:
        log_error("load_config", e)
    return False


def save_record_enabled(flag):
    """保存开关状态(失败静默, 不影响检测主流程)"""
    try:
        os.makedirs(CONFIG_DIR, exist_ok=True)
        with open(CONFIG_PATH, "w", encoding="utf-8") as f:
            f.write("[record]\nenabled = %s\n" % ("true" if flag else "false"))
    except Exception as e:
        log_error("save_config", e)


def collect_detection_rows(wmi_cache=None, edid_list=None):
    """采集待记录的显示器行: 品牌/型号/序列号/出厂年月/分辨率/刷新率"""
    wmi = wmi_cache if wmi_cache is not None else get_monitors_wmi()
    edid = edid_list if edid_list is not None else get_monitors_edid()
    edid_by = {e["mfr"].upper(): e for e in edid}
    rows = []
    for d in enum_active_displays():
        pnp = (d["pnp"] or "").upper()
        mfr = pnp[:3]
        e = edid_by.get(pnp, {})
        w = next((x for x in wmi if mfr and x["mfr"].upper().startswith(mfr)),
                 {})
        brand = MFR_MAP.get(mfr) or mfr or "未知品牌"
        model = e.get("name") or w.get("name") or ""
        serial = (e.get("serial") or "").strip()
        if not serial:
            s = (w.get("serial") or "").strip()
            serial = s if s.strip("0 ") else ""     # WMI 的 "0" 视为无效
        ym = e.get("ym") or ""
        if not ym and w.get("year") and w["year"] != "0":
            ym = w["year"] + "年"
        rows.append({
            "key": pnp or "%dx%d" % (d["w"], d["h"]),
            "brand": brand,
            "model": model or "(未知型号)",
            "serial": serial or "—",
            "ym": ym or "—",
            "res": ("%dx%d" % (d["w"], d["h"])) if d["w"] else "—",
            "hz": ("%dHz" % d["hz"]) if d["hz"] else "—",
        })
    return rows


def _write_rows(path, rows, ts):
    if os.path.exists(path):
        wb = load_workbook(path)
        ws = wb.active
        if ws.max_row == 0 or ws.cell(1, 1).value is None:
            ws.append(HEADERS)
            for c in ws[1]:
                c.font = Font(bold=True)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "检测记录"
        ws.append(HEADERS)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center")
        for i, wd in enumerate(COL_WIDTHS, 1):
            ws.column_dimensions[get_column_letter(i)].width = wd
        ws.freeze_panes = "A2"
    for r in rows:
        ws.append([ts, r["brand"], r["model"], r["serial"], r["ym"],
                   r["res"], r["hz"]])
    wb.save(path)
    return path


def append_records(rows, path=None):
    """追加写入 xlsx; 文件不存在则新建(带表头)。返回实际写入路径。
    文件被 Excel 占用时自动改写到带时间戳的副本。"""
    if not _HAS_OPENPYXL:
        raise RuntimeError("openpyxl 不可用")
    path = path or report_path()
    ts = time.strftime("%Y-%m-%d %H:%M:%S")
    try:
        return _write_rows(path, rows, ts)
    except PermissionError:
        root, ext = os.path.splitext(path)
        alt = "%s_%s%s" % (root, time.strftime("%Y%m%d_%H%M%S"), ext)
        # 原文件被占用(Excel 打开中) → 写到副本, 不丢数据
        return _write_rows(alt, rows, ts)


# ---------------------------------------------------------------- 测试画面
# (名称, 绘制函数) —— 绘制函数签名: draw(canvas, w, h)
def _solid(color):
    def d(c, w, h):
        c.configure(bg=color)
    return d


def _bars(colors, vertical=True):
    def d(c, w, h):
        c.configure(bg="#000000")
        n = len(colors)
        if vertical:
            bw = w / n
            for i, col in enumerate(colors):
                c.create_rectangle(i * bw, 0, (i + 1) * bw, h,
                                   fill=col, width=0)
        else:
            bh = h / n
            for i, col in enumerate(colors):
                c.create_rectangle(0, i * bh, w, (i + 1) * bh,
                                   fill=col, width=0)
    return d


GRAY8 = ["#%02x%02x%02x" % (v, v, v) for v in
         (0, 36, 73, 109, 146, 182, 219, 255)]
PRIM = ["#ff0000", "#00ff00", "#0000ff",
        "#ff00ff", "#00ffff", "#ffff00", "#ffffff", "#000000"]
COLORBARS = ["#ffffff", "#ffff00", "#00ffff", "#00ff00",
             "#ff00ff", "#ff0000", "#0000ff", "#000000"]
SMPTE_TOP = ["#c0c0c0", "#c0c000", "#00c0c0", "#00c000",
             "#c000c0", "#c00000", "#0000c0"]
SMPTE_MID = ["#0000c0", "#131313", "#c000c0", "#131313",
             "#00c0c0", "#131313", "#c0c0c0"]


def _smpte(c, w, h):
    c.configure(bg="#000000")
    t1, t2, t3 = h * 0.67, h * 0.08, h * 0.25
    bw = w / 7
    for i, col in enumerate(SMPTE_TOP):
        c.create_rectangle(i * bw, 0, (i + 1) * bw, t1, fill=col, width=0)
    for i, col in enumerate(SMPTE_MID):
        c.create_rectangle(i * bw, t1, (i + 1) * bw, t1 + t2,
                           fill=col, width=0)
    # 底部黑白组合条
    seg = [w * x for x in (0, 5 / 28, 10 / 28, 15 / 28, 20 / 28, 26 / 28, 1)]
    cols = ["#0000c0", "#ffffff", "#c000c0", "#000000", "#c0c0c0", "#131313"]
    for i in range(6):
        c.create_rectangle(seg[i], t1 + t2, seg[i + 1], h,
                           fill=cols[i], width=0)


def _gradient_h(c, w, h):
    """左黑右白水平渐变(用细条逼近)"""
    c.configure(bg="#000000")
    steps = 128
    bw = w / steps
    for i in range(steps):
        v = int(255 * i / (steps - 1))
        col = "#%02x%02x%02x" % (v, v, v)
        c.create_rectangle(i * bw, 0, (i + 1) * bw + 1, h, fill=col, width=0)


def _gradient_v(c, w, h):
    c.configure(bg="#000000")
    steps = 128
    bh = h / steps
    for i in range(steps):
        v = int(255 * i / (steps - 1))
        col = "#%02x%02x%02x" % (v, v, v)
        c.create_rectangle(0, i * bh, w, (i + 1) * bh + 1, fill=col, width=0)


def _checker(c, w, h):
    c.configure(bg="#ffffff")
    n = 8
    s = min(w, h) / n
    ox, oy = (w - n * s) / 2, (h - n * s) / 2
    for r in range(n):
        for col_i in range(n):
            if (r + col_i) % 2 == 0:
                c.create_rectangle(ox + col_i * s, oy + r * s,
                                   ox + (col_i + 1) * s, oy + (r + 1) * s,
                                   fill="#000000", width=0)


def _grid(c, w, h):
    c.configure(bg="#000000")
    step = 40
    for x in range(0, w, step):
        c.create_line(x, 0, x, h, fill="#00ff00")
    for y in range(0, h, step):
        c.create_line(0, y, w, y, fill="#00ff00")
    c.create_line(0, h // 2, w, h // 2, fill="#ff0000", width=2)
    c.create_line(w // 2, 0, w // 2, h, fill="#ff0000", width=2)


def _crosshair(c, w, h):
    c.configure(bg="#000000")
    c.create_line(0, h // 2, w, h // 2, fill="#ffffff", width=1)
    c.create_line(w // 2, 0, w // 2, h, fill="#ffffff", width=1)
    r = min(w, h) // 4
    c.create_oval(w // 2 - r, h // 2 - r, w // 2 + r, h // 2 + r,
                  outline="#ffffff")
    cx, cy = w // 2, h // 2
    for dx, dy in ((-1, 0), (1, 0), (0, -1), (0, 1)):
        c.create_line(cx + dx * r, cy + dy * r,
                      cx + dx * (r + 30), cy + dy * (r + 30),
                      fill="#ffffff", width=2)


SCREENS = [
    ("纯色 - 黑", _solid("#000000")),
    ("纯色 - 白", _solid("#ffffff")),
    ("纯色 - 红", _solid("#ff0000")),
    ("纯色 - 绿", _solid("#00ff00")),
    ("纯色 - 蓝", _solid("#0000ff")),
    ("纯色 - 青", _solid("#00ffff")),
    ("纯色 - 品红", _solid("#ff00ff")),
    ("纯色 - 黄", _solid("#ffff00")),
    ("灰阶 8 级", _bars(GRAY8)),
    ("彩条(竖)", _bars(COLORBARS, vertical=True)),
    ("彩条(横)", _bars(COLORBARS, vertical=False)),
    ("SMPTE 彩条", _smpte),
    ("水平渐变", _gradient_h),
    ("垂直渐变", _gradient_v),
    ("棋盘格", _checker),
    ("网格与中线", _grid),
    ("十字准星", _crosshair),
    ("红绿蓝交替(横)", _bars(["#ff0000", "#00ff00", "#0000ff"] * 3, False)),
    ("暗点检测(1%)", None),  # 特殊: 动态绘制
]


def _dark_dots(c, w, h):
    c.configure(bg="#ffffff")
    n = 24
    sw, sh = w / n, h / n
    for r in range(n):
        for col_i in range(n):
            c.create_oval(col_i * sw + sw / 2 - 2, r * sh + sh / 2 - 2,
                          col_i * sw + sw / 2 + 2, r * sh + sh / 2 + 2,
                          fill="#000000", width=0)


SCREENS[-1] = ("暗点检测(1%)", _dark_dots)


# ---------------------------------------------------------------- 主应用
def _enable_dpi_awareness():
    """必须在创建 Tk 窗口之前调用, 否则窗口按 DPI 虚拟化尺寸创建(只占屏幕1/4)"""
    try:
        ctypes.windll.shcore.SetProcessDpiAwareness(2)  # PER_MONITOR_AWARE
    except Exception:
        try:
            ctypes.windll.user32.SetProcessDPIAware()
        except Exception:
            pass


class App:
    def __init__(self, root):
        self.root = root
        self.idx = -1            # -1 = 信息页
        self.paused = False
        self._job = None
        self._refreshing = False
        self._wmi_cache = []     # WMI 结果缓存(启动/拓扑变化时更新)
        self._edid_map = {}      # EDID 面板名映射
        self._gpu_lines = []     # 显卡名(缓存, 不参与 10 秒轮询)
        self._display_ids = ()   # 活动显示器指纹, 用于检测热插拔
        self._edid_list = []     # EDID 明细(含序列号/出厂年月)
        self._recorded = set()   # 本次运行已记录的显示器, 避免重复写
        self._save_note = ""     # 记录结果提示(显示在信息栏)
        self.record_enabled = load_record_enabled()   # 检测记录: 可选, 默认关
        self._hint_job = None    # 屏幕底部临时提示的定时器

        root.title(APP_TITLE)
        root.attributes("-fullscreen", True)
        root.attributes("-topmost", True)   # 检测画面必须盖住一切(含任务栏)
        root.configure(bg="#000000", cursor="dot")

        self.canvas = tk.Canvas(root, bg="#000000",
                                highlightthickness=0, bd=0)
        self.canvas.pack(fill=tk.BOTH, expand=True)

        # 绑定
        root.bind("<Escape>", lambda e: root.destroy())
        root.bind("<space>", lambda e: self.next())
        root.bind("<Left>", lambda e: self.prev())
        root.bind("<Right>", lambda e: self.next())
        root.bind("<p>", lambda e: self.toggle_pause())
        root.bind("<P>", lambda e: self.toggle_pause())
        root.bind("<r>", lambda e: self.toggle_record())
        root.bind("<R>", lambda e: self.toggle_record())
        root.bind("<Button-1>", lambda e: self.next())            # 左键: 下一画面
        root.bind("<Button-3>", lambda e: self._right_click_once())
        root.bind("<Double-Button-3>", lambda e: root.destroy())  # 右键双击: 退出
        root.bind("<Configure>", self._on_resize)
        root.protocol("WM_DELETE_WINDOW", root.destroy)

        self.show_info()
        # 异步加载显示器信息; 信息页等待用户按键/点击确认后才开始测试
        self._load_info_async()
        self.root.after(REFRESH_MS, self._refresh_loop)

    # ---------- 信息页 ----------
    def show_info(self):
        self.canvas.delete("all")
        self.canvas.configure(bg="#101418")
        w = self.canvas.winfo_width() or self.root.winfo_screenwidth()
        h = self.canvas.winfo_height() or self.root.winfo_screenheight()
        lines = [APP_TITLE, ""]
        if not self._mon_lines_loaded:
            lines.append("正在读取显示器信息...")
        else:
            lines += self._mon_lines + [""] + self._sys_lines
        lines += ["", "空格/左键 = 开始测试    右键双击 / ESC = 退出"]
        if self.record_enabled:
            lines.append("检测记录: 已开启 → %s  (按 R 关闭)" % REPORT_NAME)
        else:
            lines.append("检测记录: 已关闭  (按 R 开启, 检测结果可存为 Excel)")
        if self._save_note:
            lines.append(self._save_note)
        # 行高按字体实际度量计算(高 DPI 下固定像素会挤行), 再整体垂直居中
        f_title = tkfont.Font(family="Microsoft YaHei UI", size=22)
        f_body = tkfont.Font(family="Microsoft YaHei UI", size=14)
        line_h = ([f_title.metrics("linespace") + 26]
                  + [f_body.metrics("linespace") + 12] * (len(lines) - 1))
        total = sum(line_h)
        y = max(40, (h - total) // 2)
        for i, t in enumerate(lines):
            big = (i == 0)
            self.canvas.create_text(
                w // 2, y, text=t, anchor="n",
                fill="#e8e8e8" if not big else "#ffffff",
                font=f_title if big else f_body,
            )
            y += line_h[i]
        self.root.title(APP_TITLE)

    _mon_lines_loaded = False
    _mon_lines = []
    _sys_lines = []
    _overlay_lines = []

    def _collect(self, full):
        """采集信息。full=True 走 WMI+显卡(开子进程, 较慢); 否则只做轻量枚举"""
        if full:
            self._wmi_cache = get_monitors_wmi()
            self._edid_list = get_monitors_edid()
            self._edid_map = {e["mfr"].upper(): e["name"]
                              for e in self._edid_list if e["name"]}
            self._gpu_lines = get_gpu_names()
            mons = merge_monitors()
            self._mon_lines = ["显示器信息:"]
            if mons:
                for m in mons:
                    self._mon_lines.append("  ● " + pretty_line(m))
            else:
                self._mon_lines.append("  (未读取到显示器信息)")
            self._sys_lines = self._gpu_lines
        mon_lines, ids = quick_monitor_lines(self._wmi_cache, self._edid_map)
        topology_changed = (ids != self._display_ids)
        self._display_ids = ids
        self._overlay_lines = (mon_lines
                               + self._gpu_lines
                               + [self._stamp()])
        self._mon_lines_loaded = True
        return topology_changed

    @staticmethod
    def _stamp():
        # 注意: 不能用 "%d" 再套 % 格式化, strftime 会先把 %d 当日期吃掉
        return (time.strftime("信息自动刷新 %H:%M:%S")
                + f" · 每 {REFRESH_MS // 1000} 秒一次")

    def _load_info_async(self, full=True):
        def worker():
            try:
                changed = self._collect(full=full)
                if changed and not full:
                    # 显示器增减/换接 → 重新走一次完整 WMI 采集
                    self._collect(full=True)
                self._maybe_record()
            except Exception as e:
                # 不再静默吞异常: 记录到 %TEMP%\MonitorCheck_err.log
                log_error("collect(full=%s)" % full, e)
            finally:
                self._refreshing = False
            if self.root.winfo_exists():
                self.root.after(0, self._apply_info_update)
        threading.Thread(target=worker, daemon=True).start()

    def _maybe_record(self):
        """把本次检测到的显示器写入 xlsx(同一台只写一次), 文件不存在则新建"""
        if not self.record_enabled:      # 记录是可选项, 未开启则完全不动磁盘
            return
        try:
            rows = collect_detection_rows(self._wmi_cache, self._edid_list)
            new = [r for r in rows if r["key"] not in self._recorded]
            if not new:
                return
            try:
                path = append_records(new)
            except Exception as e:
                log_error("record primary path", e)
                path = append_records(new, fallback_report_path())
            for r in new:
                self._recorded.add(r["key"])
            self._save_note = "已记录 %d 台显示器 → %s · %s" % (
                len(new), os.path.basename(path), time.strftime("%H:%M:%S"))
        except Exception as e:
            log_error("record_detection", e)
            self._save_note = "检测记录写入失败, 详见 %%TEMP%%\\MonitorCheck_err.log"

    # ---------- 交互 ----------
    def _right_click_once(self):
        """右键单击只提示, 必须双击才退出(防止测试中误触)"""
        self._show_hint("再按一次右键 = 退出程序")

    def toggle_record(self):
        """切换"保存检测记录"开关, 并记住选择"""
        self.record_enabled = not self.record_enabled
        save_record_enabled(self.record_enabled)
        self._save_note = ""
        if self.record_enabled:
            self._show_hint("检测记录: 已开启, 正在写入 Excel...")
        else:
            self._show_hint("检测记录: 已关闭")
        if self.idx == -1:
            self.show_info()
        else:
            self._draw_overlay()
        if self.record_enabled:
            threading.Thread(target=self._record_now, daemon=True).start()

    def _record_now(self):
        """开启记录后立即补记当前显示器(不阻塞界面)"""
        try:
            self._maybe_record()
            if not self._recorded:
                self._save_note = "未检测到可记录的显示器"
        except Exception as e:
            log_error("record_now", e)
        if self.root.winfo_exists():
            self.root.after(0, self._apply_info_update)

    def _show_hint(self, text, ms=1800):
        """屏幕底部临时提示条(独立图层, 不挤动信息栏)"""
        c = self.canvas
        self._clear_hint()
        w, h = c.winfo_width(), c.winfo_height()
        if w < 80 or h < 80:
            return
        font = tkfont.Font(family="Microsoft YaHei UI", size=13)
        tw, th = font.measure(text), font.metrics("linespace")
        pad = 14
        y0 = h - 96
        c.create_rectangle(w / 2 - tw / 2 - pad, y0 - pad,
                           w / 2 + tw / 2 + pad, y0 + th + pad,
                           fill="#101418", outline="#c8a23c", width=2,
                           tags=("hint",))
        c.create_text(w / 2, y0, text=text, anchor="n",
                      fill="#e8c86a", font=font, tags=("hint",))
        self._hint_job = self.root.after(ms, self._clear_hint)

    def _clear_hint(self):
        if self._hint_job:
            try:
                self.root.after_cancel(self._hint_job)
            except Exception:
                pass
            self._hint_job = None
        try:
            self.canvas.delete("hint")
        except Exception:
            pass

    def _apply_info_update(self):
        """信息刷新后只重绘信息层: 测试画面不重绘(不闪烁)、循环计时不重置"""
        if not self.root.winfo_exists():
            return
        if self.idx == -1:
            self.show_info()
        else:
            self._draw_overlay()

    def _refresh_loop(self):
        """每 10 秒轻量刷新一次显示器信息(纯 ctypes 枚举, 不起子进程)"""
        if not self.root.winfo_exists():
            return
        if not self._refreshing:
            self._refreshing = True
            self._load_info_async(full=False)
        self.root.after(REFRESH_MS, self._refresh_loop)

    def _draw_overlay(self):
        """测试画面中上部常驻信息栏 (深色底板保证任何背景下可读)"""
        c = self.canvas
        c.delete("overlay")          # 只清信息层, 测试画面保持原样
        if not self._mon_lines_loaded or not self._overlay_lines:
            return
        lines = list(self._overlay_lines)
        if self._save_note:
            lines.append(self._save_note)
        w = c.winfo_width()
        pad = 14
        font = tkfont.Font(family="Microsoft YaHei UI", size=10)
        lw = max(font.measure(t) for t in lines)
        lh = font.metrics("linespace")
        top = 24
        pw, ph = lw + pad * 2, len(lines) * lh + pad * 2
        x0 = (w - pw) / 2
        c.create_rectangle(x0, top, x0 + pw, top + ph, tags=("overlay",),
                           fill="#101418", outline="#4a545c")
        y = top + pad
        for t in lines:
            c.create_text(w / 2, y, text=t, anchor="n", tags=("overlay",),
                          fill="#d8d8d8" if t != self._save_note else "#8fd18f",
                          font=font)
            y += lh

    # ---------- 画面切换 ----------
    def _cancel(self):
        if self._job:
            self.root.after_cancel(self._job)
            self._job = None

    def render(self):
        self._cancel()
        name, fn = SCREENS[self.idx]
        self.canvas.delete("all")
        fn(self.canvas, self.canvas.winfo_width(),
           self.canvas.winfo_height())
        self._draw_overlay()
        self.root.title(f"{APP_TITLE}  [{self.idx + 1}/{len(SCREENS)}] {name}"
                        + ("  [已暂停]" if self.paused else "  [自动循环]"))
        if not self.paused:
            self._job = self.root.after(AUTO_MS, self.next)

    def next(self):
        self.idx = (self.idx + 1) % len(SCREENS)
        self.render()

    def prev(self):
        if self.idx == -1:      # 信息页没有上一页
            return
        self.idx = (self.idx - 1) % len(SCREENS)
        self.render()

    def toggle_pause(self):
        self.paused = not self.paused
        if self.idx >= 0:
            self.render()

    def _on_resize(self, e):
        if self.idx >= 0:
            # 防抖
            self._cancel()
            self._job = self.root.after(120, self.render)


def main():
    _enable_dpi_awareness()
    # 阻止系统休眠与熄屏: ES_CONTINUOUS | ES_SYSTEM_REQUIRED | ES_DISPLAY_REQUIRED
    _ES_CONT = 0x80000000
    ctypes.windll.kernel32.SetThreadExecutionState(
        _ES_CONT | 0x00000001 | 0x00000002)
    try:
        root = tk.Tk()
        App(root)
        root.mainloop()
    finally:
        # 恢复系统电源策略
        ctypes.windll.kernel32.SetThreadExecutionState(_ES_CONT)


if __name__ == "__main__":
    main()
